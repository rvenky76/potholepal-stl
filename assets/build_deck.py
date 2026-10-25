#!/usr/bin/env python3
"""HazardHub CAC deck. uv run --with python-pptx --with pillow python assets/build_deck.py
Output: demo/HazardHub-CAC-deck.pptx"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIA=os.path.join(ROOT,'demo','diagrams'); SHOT=os.path.join(ROOT,'demo')
OUT=os.path.join(ROOT,'demo','HazardHub-CAC-deck.pptx')
BG=RGBColor(0x15,0x20,0x2B); ACCENT=RGBColor(0xF2,0xA3,0x3C); TEXT=RGBColor(0xF1,0xF6,0xFB); MUTED=RGBColor(0x9B,0xB0,0xC4)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]; W=prs.slide_width

def bg(s): s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
def tb(s,x,y,w,h,text,size,color=TEXT,bold=False,align=PP_ALIGN.LEFT,ls=1.0):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.word_wrap=True
    for i,ln in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.line_spacing=ls
        r=p.add_run(); r.text=ln; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name='Arial'
    return box
def bar(s,x=0.0,y=0.0,w=13.333,h=0.16,color=ACCENT):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=color; sh.line.fill.background()
def pic(s,path,x,y,w,h):
    iw,ih=Image.open(path).size; ar=iw/ih
    nw,nh=(h*ar,h) if w/h>ar else (w,w/ar)
    s.shapes.add_picture(path,Inches(x+(w-nw)/2),Inches(y+(h-nh)/2),Inches(nw),Inches(nh))

# 1 Title
s=prs.slides.add_slide(BLANK); bg(s); bar(s,y=0,h=0.2)
p=os.path.join(SHOT,'screen-report.png')
if os.path.exists(p): pic(s,p,9.3,1.0,3.3,5.5)
tb(s,0.9,2.0,8.0,1.2,'HazardHub STL',56,TEXT,bold=True)
tb(s,0.9,3.2,8.0,0.8,'Report any St. Louis road hazard in 15 seconds',23,ACCENT)
tb(s,0.9,4.1,8.0,1.4,'Snap a photo, pin the spot, and HazardHub routes it to the right\nagency with a report ready to send.',20,MUTED,ls=1.2)
tb(s,0.9,6.7,9,0.4,'Congressional App Challenge  ·  React Native / Expo / TypeScript',14,MUTED)

# 2 Problem
s=prs.slides.add_slide(BLANK); bg(s)
tb(s,0.9,0.6,11.5,0.9,'The problem',40,TEXT,bold=True); bar(s,x=0.9,y=1.5,w=2.2,h=0.06)
tb(s,0.9,1.9,11.5,4.6,
   'Potholes wreck tires and cause crashes — and in St. Louis, reporting one is\n'
   'weirdly hard. Road responsibility is split across THREE agencies by who owns\n'
   'the road, and none of them has a real mobile app:\n\n'
   '• MoDOT (state highways) — a long web form with a CAPTCHA, or a phone line\n'
   '• City of St. Louis — Citizens\u2019 Service Bureau web form + phone\n'
   '• St. Louis County — a separate phone line\n\n'
   'So a resident has to guess which agency owns the road, find the form, and\n'
   'describe a location from memory. Most people just don\u2019t — and the pothole stays.\n'
   '(Same friction for a downed tree, flooding, a dead deer, or a dark signal.)',
   22,MUTED,ls=1.25)

# 3 flow diagram
s=prs.slides.add_slide(BLANK); bg(s); pic(s,os.path.join(DIA,'01-flow.png'),0.5,0.4,12.333,6.7)
# 4 screens
s=prs.slides.add_slide(BLANK); bg(s)
tb(s,0.9,0.5,11.5,0.9,'What it does',40,TEXT,bold=True); bar(s,x=0.9,y=1.4,w=2.2,h=0.06)
tb(s,0.9,1.15,11.5,0.5,'14 hazard & incident types  ·  one-tap routing  ·  live hazard map',16,ACCENT)
shots=[('screen-report.png','14 types · capture · route'),('screen-settings.png','Theme · how it works')]
x=2.3
for fn,cap in shots:
    pp=os.path.join(SHOT,fn)
    if os.path.exists(pp): pic(s,pp,x,1.8,3.6,4.5); tb(s,x,6.35,3.6,0.5,cap,15,ACCENT,bold=True,align=PP_ALIGN.CENTER)
    x+=4.6
# 5 breadth
s=prs.slides.add_slide(BLANK); bg(s)
tb(s,0.9,0.6,11.5,0.9,'More than potholes',40,TEXT,bold=True); bar(s,x=0.9,y=1.5,w=2.2,h=0.06)
tb(s,0.9,1.9,6.2,4.8,
   'One capture \u2192 locate \u2192 route \u2192 compose flow,\nnow across 14 types:\n\n'
   '\u2022 Maintenance hazards: potholes, trees,\n   flooding, dead animals, signal/light,\n   signs, debris, sidewalk, storm/ice\n'
   '\u2022 Incidents: accidents, disabled vehicles,\n   road closures, downed power lines\n\n'
   'Life-safety cases show a warning + one-tap\nCall 911 \u2014 the app says plainly it does\nNOT dispatch help itself.',
   20,MUTED,ls=1.2)
mp=os.path.join(DIA,'04-map.png')
if os.path.exists(mp): pic(s,mp,7.1,1.9,5.9,4.8)
# 6 honest  7 arch
for fn in ['02-honest.png','03-arch.png']:
    s=prs.slides.add_slide(BLANK); bg(s); pic(s,os.path.join(DIA,fn),0.5,0.4,12.333,6.7)
# 8 Close
s=prs.slides.add_slide(BLANK); bg(s); bar(s,y=0,h=0.2)
tb(s,0.9,1.6,11.5,1.0,'Make reporting so easy people actually do it',34,TEXT,bold=True)
tb(s,0.9,2.9,11.5,2.8,
   'A real, installable app that turns a multi-step scavenger hunt into a few taps \u2014\n'
   'across 14 hazard and incident types, with a map of what\u2019s been reported.\n\n'
   'Next: on-device AI to suggest the hazard type from the photo, a roads dataset for\n'
   'exact per-street routing, and a shared backend so the map goes community-wide\n'
   'with a municipal dashboard for faster agency response.',
   21,MUTED,ls=1.25)
tb(s,0.9,6.5,11.5,0.5,'A community reporting assistant. Not affiliated with MoDOT or the City/County.',15,ACCENT)

prs.save(OUT); print('Wrote',OUT,'-',len(prs.slides._sldIdLst),'slides')
